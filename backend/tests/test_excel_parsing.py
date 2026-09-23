"""Excel parsing tests (specification §6, §19)."""

from __future__ import annotations

import io
from collections.abc import Sequence
from datetime import datetime

import pytest
from openpyxl import Workbook

from app.domain.excel import ExcelParseError, TooManyRowsError, parse_excel


def make_xlsx(rows: Sequence[Sequence[object]]) -> bytes:
    """Build a real .xlsx in memory, so the parser is exercised against the actual
    format rather than a mock."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class TestParsing:
    def test_headers_and_rows(self) -> None:
        content = make_xlsx([["Phone", "Name"], ["9876543210", "Rahul"]])
        sheet = parse_excel(content, max_rows=100)

        assert sheet.headers == ["Phone", "Name"]
        assert sheet.rows == [{"Phone": "9876543210", "Name": "Rahul"}]

    def test_header_whitespace_is_trimmed(self) -> None:
        sheet = parse_excel(make_xlsx([["  Phone  "], ["9876543210"]]), max_rows=100)
        assert sheet.headers == ["Phone"]

    def test_numeric_cells_keep_their_type(self) -> None:
        """Normalization handles the int/float shapes; the parser must not stringify
        and lose that information."""
        sheet = parse_excel(make_xlsx([["Phone"], [919876543210]]), max_rows=100)
        assert sheet.rows[0]["Phone"] == 919876543210

    def test_blank_rows_are_skipped(self) -> None:
        """Trailing blanks are an editing artifact, not data, and must not be
        reported as invalid rows."""
        content = make_xlsx([["Phone"], ["9876543210"], [None], ["", None], ["9876543211"]])
        sheet = parse_excel(content, max_rows=100)
        assert len(sheet.rows) == 2

    def test_a_sheet_shorter_than_the_scan_keeps_row_one(self) -> None:
        """A file with fewer rows than the scan depth ends the search: the header is
        whatever the top row carried, exactly as the platform has always read it."""
        content = make_xlsx([["Email", "Company"], ["r@example.com", "Acme"]])

        sheet = parse_excel(content, max_rows=100)
        scanned = parse_excel(
            content,
            max_rows=100,
            header_matches=lambda headers: False,
            header_scan_rows=5,
        )

        assert scanned.headers == sheet.headers == ["Email", "Company"]
        assert scanned.header_row == sheet.header_row == 1
        assert scanned.rows == sheet.rows == [{"Email": "r@example.com", "Company": "Acme"}]

    def test_short_rows_are_padded(self) -> None:
        """Spreadsheets routinely omit trailing cells."""
        content = make_xlsx([["Phone", "Name", "Email"], ["9876543210"]])
        sheet = parse_excel(content, max_rows=100)
        assert sheet.rows[0] == {"Phone": "9876543210", "Name": None, "Email": None}

    def test_datetime_cells_survive(self) -> None:
        """Excel cannot store a timezone, so date cells arrive naive. They are kept
        rather than dropped; the repository stringifies them for JSONB storage."""
        naive = datetime(2026, 1, 1, 9, 30)  # noqa: DTZ001 - Excel dates carry no zone
        content = make_xlsx([["Phone", "Joined"], ["9876543210", naive]])
        sheet = parse_excel(content, max_rows=100)
        assert sheet.rows[0]["Joined"] == naive

    def test_only_the_first_worksheet_is_read(self) -> None:
        workbook = Workbook()
        first = workbook.active
        assert first is not None
        first.append(["Phone"])
        first.append(["9876543210"])
        second = workbook.create_sheet("Other")
        second.append(["Phone"])
        second.append(["9999999999"])

        buffer = io.BytesIO()
        workbook.save(buffer)
        sheet = parse_excel(buffer.getvalue(), max_rows=100)

        assert len(sheet.rows) == 1


class TestHeaderRowDetection:
    """A title above the headers is normal in the client's spreadsheets ([D63]).

    Academicae's real file (2026-09-24) is the shape below: a merged title, a blank row,
    then the headers. Read from row 1 it appeared to have no columns at all, and the
    upload was refused with "Missing required column(s): School Name, Phone" while the
    columns were plainly there. Only the header row's *position* is inferred, and only
    when the caller supplies the predicate.
    """

    TITLE = "Master List — All 100 Delhi Schools"

    def client_file(self) -> bytes:
        return make_xlsx(
            [
                [self.TITLE],
                [],
                ["S.No.", "Category", "School Name", "Address", "Phone", "Email"],
                [1, "referral", "Cambridge School", "New Friends Colony", 917467865887, "a@b.in"],
                [2, "referral", "Ardee School", "New Friends Colony", 919711000498, None],
            ]
        )

    @staticmethod
    def academicae(headers: Sequence[str]) -> bool:
        wanted = {"school name", "phone"}
        return wanted <= {header.strip().casefold() for header in headers}

    def test_the_header_row_is_found_below_a_title(self) -> None:
        sheet = parse_excel(
            self.client_file(), max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )

        assert sheet.header_row == 3
        assert sheet.headers[:3] == ["S.No.", "Category", "School Name"]
        assert len(sheet.rows) == 2

    def test_the_title_above_the_headers_is_not_read_as_data(self) -> None:
        """Every preamble row would otherwise be an invalid row in the preview."""
        sheet = parse_excel(
            self.client_file(), max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )

        assert sheet.rows[0]["School Name"] == "Cambridge School"
        assert all(self.TITLE not in str(value) for row in sheet.rows for value in row.values())

    def test_the_sheet_is_named_in_the_result(self) -> None:
        """An operator told a column is missing needs to know which sheet was read."""
        sheet = parse_excel(
            self.client_file(), max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )
        assert sheet.sheet_name == "Sheet"

    def test_without_a_predicate_row_one_is_still_the_header_row(self) -> None:
        """The long-standing behaviour, unchanged for a caller that states no format."""
        sheet = parse_excel(self.client_file(), max_rows=100)

        assert sheet.header_row == 1
        assert sheet.headers == [self.TITLE]

    def test_a_header_deeper_than_the_scan_reports_row_one(self) -> None:
        """Bounded on purpose: a header row further down is not searched for, and the
        caller reports the file as missing its columns rather than being guessed at."""
        content = make_xlsx(
            [[self.TITLE], [], ["note"], ["x"], [], [self.TITLE], ["School Name", "Phone"]]
        )

        sheet = parse_excel(
            content, max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )

        assert sheet.header_row == 1
        assert sheet.headers == [self.TITLE]

    def test_a_matching_header_row_is_preferred_even_when_row_one_matches(self) -> None:
        content = make_xlsx([["School Name", "Phone"], ["A", "9876543210"]])

        sheet = parse_excel(
            content, max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )

        assert sheet.header_row == 1
        assert len(sheet.rows) == 1

    def test_the_scan_never_treats_a_blank_row_as_the_header(self) -> None:
        content = make_xlsx([[self.TITLE], [], [], ["School Name", "Phone"], ["A", "9876543210"]])

        sheet = parse_excel(
            content, max_rows=100, header_matches=self.academicae, header_scan_rows=5
        )

        assert sheet.header_row == 4
        assert len(sheet.rows) == 1


class TestLimitsAndFailures:
    def test_row_limit_is_enforced(self) -> None:
        """§19: limit uploaded files. Enforced while streaming, so an oversized file
        is refused rather than exhausting memory first."""
        content = make_xlsx([["Phone"], *[[f"98765432{i:02d}"] for i in range(10)]])
        with pytest.raises(TooManyRowsError):
            parse_excel(content, max_rows=5)

    def test_a_file_at_the_limit_is_accepted(self) -> None:
        content = make_xlsx([["Phone"], *[[f"98765432{i:02d}"] for i in range(5)]])
        assert len(parse_excel(content, max_rows=5).rows) == 5

    def test_non_xlsx_content_is_rejected(self) -> None:
        with pytest.raises(ExcelParseError):
            parse_excel(b"this is not a spreadsheet", max_rows=100)

    def test_empty_bytes_are_rejected(self) -> None:
        with pytest.raises(ExcelParseError):
            parse_excel(b"", max_rows=100)

    def test_empty_worksheet_is_rejected(self) -> None:
        with pytest.raises(ExcelParseError, match="empty"):
            parse_excel(make_xlsx([]), max_rows=100)

    def test_headers_without_data_parse_to_no_rows(self) -> None:
        """Distinguishable from a broken file: the caller decides what to do."""
        sheet = parse_excel(make_xlsx([["Phone", "Name"]]), max_rows=100)
        assert sheet.headers == ["Phone", "Name"]
        assert sheet.rows == []
