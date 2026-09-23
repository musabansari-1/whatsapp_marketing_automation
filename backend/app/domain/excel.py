"""Excel parsing (specification §6, [D25], [D51], [D63]).

Turns uploaded bytes into headers and row dictionaries. Nothing here validates: that
is `app.domain.audience`, so parsing failures stay distinguishable from data problems.

§19 requires limiting uploaded files. The row cap is enforced while streaming rather
than after loading, so a file with a million rows is refused instead of exhausting
memory first.

**Where the header row is.** Row 1 by default. Client spreadsheets routinely carry a
title above the headers (often merged, often followed by a blank row), so a caller that
knows what *its* headers look like passes `header_matches` and the first matching row
within `header_scan_rows` becomes the header row. That predicate belongs to the caller
because headers are a product's configuration ([D25], [D51]) and this module is
deliberately schema-blind. Only the header's **position** is ever inferred: which
columns exist is never guessed, and a file with no matching row is read from row 1 and
reported as missing its required columns, exactly as before.
"""

from __future__ import annotations

import io
import itertools
from collections.abc import Callable
from dataclasses import dataclass

from openpyxl import load_workbook


class ExcelParseError(Exception):
    """The upload is not a readable spreadsheet."""


class TooManyRowsError(ExcelParseError):
    """The file exceeds the configured row limit (§19)."""

    def __init__(self, limit: int) -> None:
        super().__init__(f"File exceeds the maximum of {limit} rows")
        self.limit = limit


@dataclass(frozen=True, slots=True)
class ParsedSheet:
    """One worksheet, read.

    `header_row` and `sheet_name` are carried so a caller can say *where* the headers
    came from. When a file is reported as missing columns, which row and which sheet
    were read is the first thing an operator needs and the last thing they can guess.
    """

    headers: list[str]
    rows: list[dict[str, object]]
    header_row: int = 1
    sheet_name: str = ""


def parse_excel(
    content: bytes,
    *,
    max_rows: int,
    header_matches: Callable[[list[str]], bool] | None = None,
    header_scan_rows: int = 1,
) -> ParsedSheet:
    """Parse the first worksheet into headers and row dictionaries.

    `read_only` and `data_only` matter: the first streams rather than building the
    whole document in memory, and the second yields cached formula *results* rather
    than formula text, since a cell containing `=A1` should contribute its value.
    """
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
    except Exception as exc:  # noqa: BLE001 - openpyxl raises a variety of types
        raise ExcelParseError("The file could not be read as an .xlsx workbook") from exc

    try:
        worksheet = workbook.worksheets[0] if workbook.worksheets else None
        if worksheet is None:
            raise ExcelParseError("The workbook contains no worksheets")

        # Without a predicate there is nothing to look for, so row 1 is the header row
        # and no other row is considered.
        limit = max(1, header_scan_rows) if header_matches is not None else 1
        row_iterator = worksheet.iter_rows(values_only=True)

        # The scan reads at most `limit` rows, and those rows are buffered because a
        # failed scan must not eat them: the fallback reads row 1 as the header and the
        # rows after it as data, which is only possible if the scan handed them back.
        scanned: list[tuple[object, ...]] = []
        headers: list[str] = []
        header_row = 1

        for row in row_iterator:
            scanned.append(row)
            candidate = _clean_headers(row)

            if candidate and header_matches is not None and header_matches(candidate):
                headers = candidate
                header_row = len(scanned)
                break

            if len(scanned) >= limit:
                break

        if not scanned:
            raise ExcelParseError("The worksheet is empty")
        if not headers:
            # Nothing in the scanned range is this format's header row. Fall back to
            # row 1 — the platform's long-standing behaviour — so the caller reports
            # missing columns rather than this module inventing a layout.
            headers = _clean_headers(scanned[0])
            if not headers:
                raise ExcelParseError("The first row must contain column headers")

        # Rows above the header are preamble (a title, a note) and are deliberately not
        # read as data: they are not contacts and every one of them would be an invalid
        # row in the preview. Whatever the scan buffered below the header is returned
        # first, then the rest of the sheet streams.
        remaining = itertools.chain(
            (values for index, values in enumerate(scanned, start=1) if index > header_row),
            row_iterator,
        )

        rows: list[dict[str, object]] = []
        for values in remaining:
            if _is_blank(values):
                # Trailing blank rows are an artifact of how spreadsheets are edited,
                # not data, and must not be reported as invalid rows.
                continue
            if len(rows) >= max_rows:
                raise TooManyRowsError(max_rows)
            rows.append(_to_dict(headers, values))

        return ParsedSheet(
            headers=headers, rows=rows, header_row=header_row, sheet_name=worksheet.title
        )
    finally:
        workbook.close()


def _clean_headers(header_row: tuple[object, ...]) -> list[str]:
    """Header texts, positionally preserved.

    Empty headers become positional placeholders rather than being dropped, so the
    columns after a blank header still line up with their values.
    """
    headers: list[str] = []
    for index, value in enumerate(header_row):
        text = "" if value is None else str(value).strip()
        headers.append(text or f"column_{index + 1}")

    while headers and headers[-1].startswith("column_"):
        headers.pop()
    return headers


def _is_blank(values: tuple[object, ...]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


def _to_dict(headers: list[str], values: tuple[object, ...]) -> dict[str, object]:
    """Zip headers with values, tolerating short and long rows.

    Spreadsheets routinely produce rows shorter than the header row; treating the
    missing cells as absent is correct, and extra cells beyond the headers are
    ignored because there is no column to attribute them to.
    """
    row: dict[str, object] = {}
    for index, header in enumerate(headers):
        row[header] = values[index] if index < len(values) else None
    return row
